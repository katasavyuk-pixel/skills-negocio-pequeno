"""The quarter review: every defect a spreadsheet can hide is found, and the totals are the sums done by hand."""
import contextlib
import datetime
import io
import json
import pathlib
import sys
import tempfile
import unittest
from decimal import Decimal

RAIZ = pathlib.Path(__file__).resolve().parent.parent
SKILL = RAIZ / "skills" / "cerrar-trimestre"
sys.path.insert(0, str(SKILL / "scripts"))

import revisar  # noqa: E402

CABECERA = "numero,fecha,nif,nombre,base,tipo_iva,cuota_iva,retencion,total\n"
D = Decimal
# Tax numbers with province code 00, which does not exist: they pass the control check and belong to nobody.
CLIENTE, OTRO, PROVEEDOR = "B00333336", "B00111112", "A00555557"


def ejecutar(emitidas=None, recibidas=None, *opciones, trimestre="2026-3"):
    carpeta = pathlib.Path(tempfile.mkdtemp())
    argumentos = ["--trimestre", trimestre]
    for nombre, contenido in (("emitidas", emitidas), ("recibidas", recibidas)):
        if contenido is not None:
            (carpeta / (nombre + ".csv")).write_text(contenido, encoding="utf-8")
            argumentos += ["--" + nombre, str(carpeta / (nombre + ".csv"))]
    salida, errores = io.StringIO(), io.StringIO()
    with contextlib.redirect_stdout(salida), contextlib.redirect_stderr(errores):
        codigo = revisar.main(argumentos + list(opciones))
    return codigo, salida.getvalue(), errores.getvalue()


def informe(emitidas=None, recibidas=None, *opciones, **claves):
    codigo, salida, errores = ejecutar(emitidas, recibidas, *opciones, **claves)
    assert codigo == 0, errores
    return json.loads(salida)


def hallazgos(datos, factura=None):
    return [h for h in datos["hallazgos"] if factura is None or h["factura"] == factura]


BUENA = "F-001,2026-07-03,%s,Óptica Lumen,2000.00,21,420.00,0,2420.00\n" % CLIENTE


class Nif(unittest.TestCase):
    VALIDOS = [
        "12345678Z",  # the usual example of a DNI
        "00000000T",
        "X1234567L",  # the usual example of a NIE
        "A58818501",
        "B12345674",
        "Q2826000H",  # the Agencia Tributaria's own number: an entity whose control character is a letter
        "ES12345678Z",
        "12.345.678-Z",
        " b12345674 ",
    ]

    def test_numeros_conocidos_como_validos(self):
        for nif in self.VALIDOS:
            with self.subTest(nif):
                self.assertEqual(revisar.comprobar_nif(nif)[0], "valido")

    def test_una_letra_cambiada_no_pasa(self):
        for nif in ("12345678A", "X1234567A", "B12345670", "Q2826000A", "A58818502"):
            with self.subTest(nif):
                self.assertEqual(revisar.comprobar_nif(nif)[0], "control")

    def test_cada_dni_tiene_una_sola_letra_valida(self):
        validas = [letra for letra in "ABCDEFGHIJKLMNOPQRSTUVWXYZ" if revisar.comprobar_nif("50123456" + letra)[0] == "valido"]
        self.assertEqual(len(validas), 1)

    def test_lo_que_no_es_un_nif_espanol(self):
        self.assertEqual(revisar.comprobar_nif("")[0], "vacio")
        self.assertEqual(revisar.comprobar_nif("1234567Z")[0], "formato")
        self.assertEqual(revisar.comprobar_nif("FR12345678901")[0], "extranjero")


class Trimestres(unittest.TestCase):
    def test_los_cuatro(self):
        self.assertEqual(revisar.trimestre_de("2026-1")[2:], (datetime.date(2026, 1, 1), datetime.date(2026, 3, 31)))
        self.assertEqual(revisar.trimestre_de("2026-2")[2:], (datetime.date(2026, 4, 1), datetime.date(2026, 6, 30)))
        self.assertEqual(revisar.trimestre_de("2026-3")[2:], (datetime.date(2026, 7, 1), datetime.date(2026, 9, 30)))
        self.assertEqual(revisar.trimestre_de("2026-4")[2:], (datetime.date(2026, 10, 1), datetime.date(2026, 12, 31)))

    def test_los_plazos_son_los_del_calendario_del_contribuyente(self):
        # 303 and 130: until the 20th of April, July and October, and 30 January. 111 and 115: until the 20th.
        self.assertEqual([p["hasta"] for p in revisar.calendario(2026, 3)], ["2026-10-20"] * 4)
        cuarto = revisar.calendario(2026, 4)
        # 30-Jan-2027 is a Saturday: it moves to Monday 1 February.
        self.assertEqual(cuarto[0]["hasta"], "2027-02-01")
        self.assertEqual(cuarto[2]["hasta"], "2027-01-20")

    def test_un_plazo_en_fin_de_semana_pasa_al_lunes(self):
        # 20-Apr-2024 was a Saturday.
        self.assertEqual(revisar.calendario(2024, 1)[0]["hasta"], "2024-04-22")

    def test_fuera_del_territorio_comun_no_se_dan_plazos(self):
        self.assertEqual(informe(CABECERA + BUENA, None, "--sin-calendario-fiscal")["plazos"], [])


class UnaFactura(unittest.TestCase):
    def test_una_factura_correcta_no_da_nada(self):
        datos = informe(CABECERA + BUENA)
        self.assertEqual((datos["errores"], datos["avisos"]), (0, 0))

    CASOS = [
        ("la cuota no es la base por el tipo", "F-001,2026-07-03,%s,Lumen,1500.00,21,300.00,0,1800.00\n" % CLIENTE, revisar.ERROR, "saldrían 315.00"),
        ("el total no es la suma", "F-001,2026-07-03,%s,Lumen,100.00,21,21.00,0,112.00\n" % CLIENTE, revisar.ERROR, "saldrían 121.00"),
        ("la retención se olvida al sumar", "F-001,2026-07-03,%s,Lumen,1000.00,21,210.00,150.00,1210.00\n" % CLIENTE, revisar.ERROR, "saldrían 1060.00"),
        ("un NIF con la letra mal", "F-001,2026-07-03,B00333330,Lumen,2000.00,21,420.00,0,2420.00\n", revisar.ERROR, "dígito de control"),
        ("sin NIF del destinatario", "F-001,2026-07-03,,Bar Central,800.00,21,168.00,0,968.00\n", revisar.ERROR, "Falta el NIF del destinatario"),
        ("sin nombre", "F-001,2026-07-03,%s,,2000.00,21,420.00,0,2420.00\n" % CLIENTE, revisar.ERROR, "Falta el nombre"),
        ("de un trimestre anterior", "F-001,2026-06-28,%s,Lumen,2000.00,21,420.00,0,2420.00\n" % CLIENTE, revisar.ERROR, "anterior al trimestre"),
        ("de un trimestre posterior", "F-001,2026-10-02,%s,Lumen,2000.00,21,420.00,0,2420.00\n" % CLIENTE, revisar.AVISO, "posterior al trimestre"),
        ("un tipo raro", "F-001,2026-07-03,%s,Lumen,100.00,5,5.00,0,105.00\n" % CLIENTE, revisar.AVISO, "no es ninguno de los habituales"),
        ("un cliente de fuera", "F-001,2026-07-03,FR12345678901,Client,100.00,0,0.00,0,100.00\n", revisar.AVISO, "fuera de España"),
    ]

    def test_cada_defecto_se_encuentra_y_se_explica(self):
        for nombre, linea, gravedad, mensaje in self.CASOS:
            with self.subTest(nombre):
                encontrados = hallazgos(informe(CABECERA + linea), "F-001")
                self.assertTrue(any(h["gravedad"] == gravedad and mensaje in h["que"] for h in encontrados), encontrados)
                self.assertTrue(all(h["arreglo"] for h in encontrados))

    def test_dos_centimos_de_redondeo_no_son_un_error(self):
        # 33.33 x 21 % = 6.9993: an invoice that rounds line by line may print 7.00 or 6.99.
        datos = informe(CABECERA + "F-001,2026-07-03,%s,Lumen,33.33,21,7.01,0,40.34\n" % CLIENTE)
        self.assertEqual(datos["errores"], 0)

    def test_una_factura_con_dos_tipos_son_dos_lineas_y_una_factura(self):
        datos = informe(CABECERA + "F-001,2026-09-15,%s,Paso,800.00,21,168.00,0,1188.00\nF-001,2026-09-15,%s,Paso,200.00,10,20.00,0,1188.00\n" % (OTRO, OTRO))
        self.assertEqual((datos["facturas_emitidas"], datos["errores"]), (1, 0))

    def test_una_rectificativa_en_negativo_cuadra(self):
        datos = informe(CABECERA + "R-001,2026-07-03,%s,Lumen,-500.00,21,-105.00,0,-605.00\n" % CLIENTE)
        self.assertEqual(datos["errores"], 0)
        self.assertEqual(datos["iva_de_las_ventas"]["total"], "-105.00")

    def test_recargo_de_equivalencia(self):
        cabecera = CABECERA.strip() + ",recargo\n"
        bien = informe(cabecera + "F-001,2026-07-03,%s,Tienda,1000.00,21,210.00,0,1262.00,52.00\n" % CLIENTE)
        self.assertEqual(bien["errores"], 0)
        self.assertEqual(bien["iva_de_las_ventas"]["total"], "262.00")
        mal = informe(cabecera + "F-001,2026-07-03,%s,Tienda,1000.00,21,210.00,0,1224.00,14.00\n" % CLIENTE)
        self.assertTrue(any("saldrían 52.00" in h["que"] for h in hallazgos(mal)))

    def test_simplificadas(self):
        cabecera = CABECERA.strip() + ",simplificada\n"
        tique = informe(cabecera + "T-001,2026-07-03,,Venta de mostrador,100.00,21,21.00,0,121.00,si\n")
        self.assertEqual((tique["errores"], tique["avisos"]), (0, 0))
        medio = informe(cabecera + "T-001,2026-07-03,,Venta,500.00,21,105.00,0,605.00,si\n")
        self.assertTrue(any(h["gravedad"] == revisar.AVISO and "art. 4.2" in h["arreglo"] for h in hallazgos(medio)))
        grande = informe(cabecera + "T-001,2026-07-03,,Venta,3000.00,21,630.00,0,3630.00,si\n")
        self.assertTrue(any(h["gravedad"] == revisar.ERROR and "factura completa" in h["arreglo"] for h in hallazgos(grande)))


class Numeracion(unittest.TestCase):
    def lineas(self, *facturas):
        return CABECERA + "".join("%s,%s,%s,Cliente,100.00,21,21.00,0,121.00\n" % (numero, fecha, nif) for numero, fecha, nif in facturas)

    def test_un_hueco_se_nombra(self):
        datos = informe(self.lineas(("2026-025", "2026-07-10", CLIENTE), ("2026-027", "2026-07-12", OTRO)))
        self.assertTrue(any("falta una factura: 2026-026" in h["que"] for h in hallazgos(datos)))

    def test_varios_huecos_seguidos(self):
        datos = informe(self.lineas(("A-8", "2026-07-10", CLIENTE), ("A-12", "2026-07-12", OTRO)))
        self.assertTrue(any("faltan 3 facturas" in h["que"] for h in hallazgos(datos)))

    def test_el_orden_de_fechas(self):
        datos = informe(self.lineas(("2026-025", "2026-07-10", CLIENTE), ("2026-026", "2026-07-01", OTRO)))
        self.assertTrue(any("fecha anterior" in h["que"] for h in hallazgos(datos)))

    def test_el_mismo_numero_a_dos_clientes(self):
        datos = informe(self.lineas(("2026-025", "2026-07-10", CLIENTE), ("2026-025", "2026-07-11", OTRO)))
        self.assertTrue(any(h["gravedad"] == revisar.ERROR and "repetido" in h["que"] for h in hallazgos(datos)))

    def test_cada_serie_se_numera_aparte(self):
        cabecera = CABECERA.strip() + ",serie\n"
        datos = informe(cabecera + "1,2026-07-10,%s,Cliente,100.00,21,21.00,0,121.00,A\n1,2026-07-11,%s,Cliente,100.00,21,21.00,0,121.00,R\n" % (CLIENTE, OTRO))
        self.assertEqual(datos["errores"], 0)

    def test_una_numeracion_correlativa_no_da_nada(self):
        datos = informe(self.lineas(("2026-025", "2026-07-10", CLIENTE), ("2026-026", "2026-07-10", OTRO), ("2026-027", "2026-08-01", CLIENTE)))
        self.assertEqual((datos["errores"], datos["avisos"]), (0, 0))


class Recibidas(unittest.TestCase):
    def test_la_misma_factura_dos_veces_se_dice_una_vez_y_no_suma_dos(self):
        linea = "H-220,2026-08-18,%s,Herrajes,1000.00,21,210.00,0,1210.00\n" % PROVEEDOR
        datos = informe(None, CABECERA + linea + linea)
        self.assertEqual(datos["errores"], 1)
        self.assertIn("ya está apuntada en la línea 2", hallazgos(datos)[0]["que"])
        self.assertEqual(datos["iva_de_las_compras"]["total"], "210.00")

    def test_un_proveedor_que_falta_el_ultimo_mes_se_pregunta(self):
        datos = informe(None, CABECERA + "A,2026-07-15,%s,Maderas,100.00,21,21.00,0,121.00\nB,2026-08-14,%s,Maderas,100.00,21,21.00,0,121.00\n" % (PROVEEDOR, PROVEEDOR))
        self.assertEqual(datos["posibles_facturas_que_faltan"], ["Maderas tiene factura dos meses del trimestre y no en septiembre. ¿Falta esa factura?"])

    def test_un_proveedor_que_empieza_el_segundo_mes_no_se_pregunta(self):
        datos = informe(None, CABECERA + "A,2026-08-15,%s,Nuevo,100.00,21,21.00,0,121.00\nB,2026-09-14,%s,Nuevo,100.00,21,21.00,0,121.00\n" % (PROVEEDOR, PROVEEDOR))
        self.assertEqual(datos["posibles_facturas_que_faltan"], [])

    def test_una_recibida_de_otro_trimestre_solo_avisa(self):
        datos = informe(None, CABECERA + "A,2026-06-20,%s,Maderas,100.00,21,21.00,0,121.00\n" % PROVEEDOR)
        self.assertEqual((datos["errores"], datos["avisos"]), (0, 1))
        self.assertEqual(datos["iva_de_las_compras"]["total"], "0.00")

    def test_las_retenciones_que_practica_el_negocio(self):
        datos = informe(None, CABECERA + "ALQ-07,2026-07-01,%s,Inmobiliaria,1198.35,21,251.65,227.69,1222.31\n" % PROVEEDOR)
        self.assertEqual(datos["errores"], 0)
        self.assertEqual(datos["retenciones_que_has_practicado"], "227.69")


class EntradaMala(unittest.TestCase):
    def test_cada_problema_del_listado_sale_con_su_libro_y_su_linea(self):
        codigo, salida, errores = ejecutar(CABECERA + "F-1,03/07/2026,%s,Lumen,\"2.000,00\",21%%,420.00,0,2420.00\n" % CLIENTE)
        self.assertEqual(codigo, 1)
        self.assertEqual(salida, "")
        self.assertIn("emitidas, línea 2: la fecha", errores)
        self.assertIn("`base` es '2.000,00'", errores)
        self.assertIn("`tipo_iva` es '21%'", errores)

    def test_trimestre_mal_escrito(self):
        codigo, _, errores = ejecutar(CABECERA + BUENA, trimestre="3T2026")
        self.assertEqual(codigo, 1)
        self.assertIn("AAAA-N", errores)

    def test_sin_ningun_listado(self):
        with self.assertRaises(SystemExit), contextlib.redirect_stderr(io.StringIO()):
            revisar.main(["--trimestre", "2026-3"])


class Ejemplo(unittest.TestCase):
    """The worked example the skill shows: the figures in its documents are these."""

    EMITIDAS = (SKILL / "ejemplos" / "emitidas.csv").read_text(encoding="utf-8")
    RECIBIDAS = (SKILL / "ejemplos" / "recibidas.csv").read_text(encoding="utf-8")

    def test_el_ejemplo_da_lo_que_dice_el_resumen_guardado(self):
        datos = informe(self.EMITIDAS, self.RECIBIDAS)
        self.assertEqual(datos, json.loads((SKILL / "ejemplos" / "resumen.json").read_text(encoding="utf-8")))

    def test_las_cifras_del_ejemplo_hechas_a_mano(self):
        datos = informe(self.EMITIDAS, self.RECIBIDAS)
        ventas_21 = D("420") + D("1260") + D("315") + D("300") + D("168") + D("630") + D("168")
        self.assertEqual(datos["iva_de_las_ventas"]["total"], str(ventas_21 + D("120")) + ".00")
        compras = D("676.86") + D("378.00") + 3 * D("251.65") + D("65.10") + D("61.95") + D("222.15") + D("21.00")
        self.assertEqual(datos["iva_de_las_compras"]["total"], str(compras))
        self.assertEqual(D(datos["diferencia"]), ventas_21 + D("120") - compras)
        self.assertEqual(datos["retenciones_que_has_practicado"], "683.07")
        self.assertEqual((datos["errores"], datos["avisos"]), (5, 4))
        self.assertEqual(len(datos["posibles_facturas_que_faltan"]), 2)

    def test_el_libro_tiene_las_cinco_hojas(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl no está instalado (en CI sí lo está).")
        libro = pathlib.Path(tempfile.mkdtemp()) / "trimestre.xlsx"
        codigo, salida, errores = ejecutar(self.EMITIDAS, self.RECIBIDAS, "--salida", str(libro))
        self.assertEqual(codigo, 0, errores)
        hojas = openpyxl.load_workbook(libro)
        self.assertEqual(hojas.sheetnames, ["Resumen", "Qué corregir", "IVA por tipo", "Emitidas", "Recibidas"])
        self.assertEqual(hojas["Qué corregir"].max_row - 1, 9)
        self.assertEqual(hojas["Emitidas"].max_row - 1, 10)


if __name__ == "__main__":
    unittest.main()
