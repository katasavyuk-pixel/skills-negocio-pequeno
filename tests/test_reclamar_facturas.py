"""Late-payment claims: every figure checked against a sum done by hand, and the legal constants against the law."""
import contextlib
import datetime
import io
import json
import pathlib
import sys
import tempfile
import unittest
from decimal import Decimal

RAIZ = pathlib.Path(__file__).resolve().parent.parent
SKILL = RAIZ / "skills" / "reclamar-facturas"
sys.path.insert(0, str(SKILL / "scripts"))

import calcular  # noqa: E402

CABECERA = "cliente,numero,fecha_factura,vencimiento,importe,tipo_cliente,avisos\n"
D = Decimal
HOY = ("--hasta", "2026-10-13")


def ejecutar(csv, *opciones):
    carpeta = pathlib.Path(tempfile.mkdtemp())
    ruta = carpeta / "facturas.csv"
    ruta.write_text(csv, encoding="utf-8")
    salida, errores = io.StringIO(), io.StringIO()
    with contextlib.redirect_stdout(salida), contextlib.redirect_stderr(errores):
        codigo = calcular.main([str(ruta), *opciones])
    return codigo, salida.getvalue(), errores.getvalue()


def resumen(csv, *opciones):
    codigo, salida, errores = ejecutar(csv, *opciones)
    assert codigo == 0, errores
    return json.loads(salida)


def factura(datos, numero):
    return next(f for f in datos["facturas"] if f["numero"] == numero)


class Intereses(unittest.TestCase):
    def test_un_solo_semestre(self):
        # Due 28-Sep-2026, calculated on 13-Oct: 15 days at 10.40 %. 3630 x 0.104 x 15 / 365 = 15.5145...
        datos = resumen(CABECERA + "Hotel,F1,2026-08-29,2026-09-28,3630.00,empresa,0\n", *HOY)
        f = factura(datos, "F1")
        self.assertEqual(f["dias_de_retraso"], 15)
        self.assertEqual(f["intereses"], "15.51")
        self.assertEqual(f["compensacion"], "40.00")
        self.assertEqual(f["total"], "3685.51")

    def test_el_cambio_de_semestre_cambia_el_tipo(self):
        # Due 9-Jun-2026: 21 days of June at 10.15 % and 105 days from 1-Jul to 13-Oct at 10.40 %.
        primer = D("7260") * D("0.1015") * 21 / 365
        segundo = D("7260") * D("0.1040") * 105 / 365
        self.assertEqual(calcular.centimos(primer + segundo), D("259.60"))
        datos = resumen(CABECERA + "Hotel,F2,2026-04-10,2026-06-09,7260.00,empresa,0\n", *HOY)
        f = factura(datos, "F2")
        self.assertEqual(f["dias_de_retraso"], 126)
        self.assertEqual(f["intereses"], "259.60")

    def test_los_tramos_cubren_cada_dia_una_sola_vez(self):
        tipos, _ = calcular.leer_tipos(calcular.TIPOS_POR_DEFECTO)
        tramos, faltan = calcular.tramos_de_interes(datetime.date(2024, 11, 15), datetime.date(2026, 10, 13), D("1000"), tipos)
        self.assertEqual(faltan, [])
        self.assertEqual([t["semestre"] for t in tramos], ["2024-2", "2025-1", "2025-2", "2026-1", "2026-2"])
        self.assertEqual(sum(t["dias"] for t in tramos), (datetime.date(2026, 10, 13) - datetime.date(2024, 11, 15)).days)
        self.assertEqual(tramos[0]["desde"], datetime.date(2024, 11, 16))
        for anterior, siguiente in zip(tramos, tramos[1:]):
            self.assertEqual(siguiente["desde"], anterior["hasta"] + datetime.timedelta(days=1))

    def test_el_dia_del_vencimiento_no_cuenta(self):
        datos = resumen(CABECERA + "Hotel,F3,2026-09-13,2026-10-13,1000.00,empresa,0\n", *HOY)
        self.assertEqual(factura(datos, "F3")["estado"], "no vencida")
        self.assertEqual(datos["total_a_reclamar"], "0.00")
        un_dia = resumen(CABECERA + "Hotel,F3,2026-09-12,2026-10-12,1000.00,empresa,0\n", *HOY)
        # One day at 10.40 %: 1000 x 0.104 / 365 = 0.2849...
        self.assertEqual(factura(un_dia, "F3")["intereses"], "0.28")

    def test_sobre_la_base_sin_iva_sale_la_cifra_prudente(self):
        cabecera = CABECERA.strip() + ",base\n"
        datos = resumen(cabecera + "Hotel,F1,2026-08-29,2026-09-28,3630.00,empresa,0,3000.00\n", *HOY)
        # 3000 x 0.104 x 15 / 365 = 12.8219...
        self.assertEqual(datos["intereses_si_se_calculan_sin_iva"], "12.82")
        self.assertEqual(datos["intereses_de_demora"], "15.51")


class LoQueDiceLaLey(unittest.TestCase):
    def test_las_constantes_son_las_de_la_ley_3_2004(self):
        self.assertEqual(calcular.PLAZO_SIN_PACTO, 30)  # art. 4.1
        self.assertEqual(calcular.PLAZO_MAXIMO, 60)  # art. 4.3
        self.assertEqual(calcular.COMPENSACION, D("40.00"))  # art. 8.1

    def test_sin_plazo_pactado_son_treinta_dias_desde_la_entrega(self):
        datos = resumen(CABECERA + "Hotel,F4,2026-08-01,,1000.00,empresa,0\n", *HOY)
        f = factura(datos, "F4")
        self.assertEqual(f["vencimiento_legal"], "2026-08-31")
        self.assertIn("art. 4.1", f["notas"][0])

    def test_un_vencimiento_a_mas_de_sesenta_dias_se_recorta(self):
        datos = resumen(CABECERA + "Zapatería,F5,2026-07-01,2026-09-29,1815.00,empresa,0\n", *HOY)
        f = factura(datos, "F5")
        self.assertEqual(f["vencimiento_legal"], "2026-08-30")
        self.assertEqual(f["dias_de_retraso"], 44)
        # 1815 x 0.104 x 44 / 365 = 22.7546...
        self.assertEqual(f["intereses"], "22.75")
        self.assertIn("art. 4.3", f["notas"][0])

    def test_la_fecha_de_entrega_manda_sobre_la_de_la_factura(self):
        cabecera = CABECERA.strip() + ",entrega\n"
        datos = resumen(cabecera + "Hotel,F6,2026-08-01,,1000.00,empresa,0,2026-08-11\n", *HOY)
        self.assertEqual(factura(datos, "F6")["vencimiento_legal"], "2026-09-10")

    def test_el_plazo_habitual_del_perfil_tambien_se_recorta_a_sesenta(self):
        datos = resumen(CABECERA + "Hotel,F7,2026-07-01,,1000.00,empresa,0\n", *HOY, "--plazo-dias", "90")
        self.assertEqual(factura(datos, "F7")["vencimiento_legal"], "2026-08-30")

    def test_a_un_particular_no_se_le_anaden_intereses_ni_compensacion(self):
        datos = resumen(CABECERA + "Familia,F8,2026-08-30,2026-09-29,2900.00,particular,1\n", *HOY)
        f = factura(datos, "F8")
        self.assertEqual((f["intereses"], f["compensacion"], f["total"]), ("0.00", "0.00", "2900.00"))
        self.assertIn("art. 3.2.a", f["notas"][0])
        self.assertEqual(datos["lo_que_la_ley_anade_al_importe"], "0.00")

    def test_una_administracion_no_se_calcula_y_se_dice(self):
        datos = resumen(CABECERA + "Ayuntamiento,F9,2026-05-04,2026-06-03,1210.00,administracion,0\n", *HOY)
        self.assertEqual(factura(datos, "F9")["total"], "0.00")
        self.assertEqual(datos["facturas_vencidas"], 0)
        self.assertTrue(any("Administración" in aviso for aviso in datos["avisos"]))

    def test_cuarenta_euros_por_factura_no_por_cliente(self):
        datos = resumen(CABECERA + "Hotel,A,2026-08-29,2026-09-28,100.00,empresa,0\nHotel,B,2026-08-29,2026-09-28,100.00,empresa,0\n", *HOY)
        self.assertEqual(datos["compensacion_por_costes_de_cobro"], "80.00")
        self.assertEqual(datos["por_cliente"][0]["compensacion"], "80.00")


class TiposDelBoe(unittest.TestCase):
    def test_la_tabla_trae_cada_tipo_con_su_resolucion(self):
        tabla = json.loads(calcular.TIPOS_POR_DEFECTO.read_text(encoding="utf-8"))
        semestres = [fila["semestre"] for fila in tabla["tipos"]]
        self.assertEqual(semestres, sorted(semestres))
        self.assertEqual(len(set(semestres)), len(semestres))
        for fila in tabla["tipos"]:
            self.assertRegex(fila["boe"], r"^BOE-A-\d{4}-\d+$")
            self.assertEqual(fila["url"], "https://www.boe.es/diario_boe/txt.php?id=" + fila["boe"])
            self.assertTrue(D("8") <= D(fila["tipo"]) <= D("20"), fila)
        datetime.date.fromisoformat(tabla["comprobado"])

    def test_no_hay_huecos_entre_semestres(self):
        tabla = json.loads(calcular.TIPOS_POR_DEFECTO.read_text(encoding="utf-8"))
        semestres = [fila["semestre"] for fila in tabla["tipos"]]
        for anterior, siguiente in zip(semestres, semestres[1:]):
            anio, mitad = (int(x) for x in anterior.split("-"))
            self.assertEqual(siguiente, "%d-2" % anio if mitad == 1 else "%d-1" % (anio + 1))

    def test_un_semestre_sin_tipo_para_el_calculo_en_vez_de_suponerlo(self):
        codigo, salida, errores = ejecutar(CABECERA + "Hotel,F1,2026-11-01,2026-12-01,1000.00,empresa,0\n", "--hasta", "2027-02-01")
        self.assertEqual(codigo, 1)
        self.assertEqual(salida, "")
        self.assertIn("FALTA EL TIPO DE DEMORA", errores)
        self.assertIn("2027-1", errores)
        self.assertIn("--tipo 2027-1=", errores)

    def test_un_tipo_aportado_se_usa_y_queda_marcado(self):
        datos = resumen(CABECERA + "Hotel,F1,2026-11-01,2026-12-31,1000.00,empresa,0\n", "--hasta", "2027-01-31", "--tipo", "2027-1=10.00")
        # 31 days at 10 %: 1000 x 0.10 x 31 / 365 = 8.4931...
        self.assertEqual(factura(datos, "F1")["intereses"], "8.49")

    def test_una_factura_anterior_a_la_tabla_tambien_para(self):
        codigo, _, errores = ejecutar(CABECERA + "Hotel,F1,2020-01-01,2020-02-01,1000.00,empresa,0\n", *HOY)
        self.assertEqual(codigo, 1)
        self.assertIn("2020-1", errores)


class Fases(unittest.TestCase):
    CASOS = [
        (10, 0, 1),
        (31, 0, 2),
        (10, 1, 2),
        (61, 0, 3),
        (10, 2, 3),
        (200, 5, 3),
    ]

    def test_la_fase_sale_de_los_avisos_enviados_y_del_tiempo(self):
        for dias, avisos, esperada in self.CASOS:
            with self.subTest(dias=dias, avisos=avisos):
                self.assertEqual(calcular.fase_de(dias, avisos), esperada)

    def test_al_cliente_le_toca_la_fase_de_su_factura_mas_avanzada(self):
        datos = resumen(CABECERA + "Hotel,A,2026-04-10,2026-06-09,100.00,empresa,2\nHotel,B,2026-08-29,2026-09-28,100.00,empresa,0\n", *HOY)
        self.assertEqual(datos["por_cliente"][0]["fase"], 3)
        self.assertEqual(datos["por_cliente"][0]["que_toca"], "requerimiento formal con propuesta de acuerdo")


class Prescripcion(unittest.TestCase):
    def test_mas_de_tres_anos_avisa_y_mas_de_cinco_avisa_mas(self):
        tres = resumen(CABECERA + "Hotel,F1,2023-01-10,2023-02-09,1000.00,empresa,0\n", *HOY)
        self.assertTrue(any("art. 1967" in nota for nota in factura(tres, "F1")["notas"]))
        cinco = resumen(CABECERA + "Hotel,F1,2021-07-10,2021-08-09,1000.00,empresa,0\n", *HOY)
        self.assertTrue(any("art. 1964.2" in nota for nota in factura(cinco, "F1")["notas"]))
        self.assertTrue(any("prescrito" in aviso for aviso in cinco["avisos"]))


class EntradaMala(unittest.TestCase):
    CASOS = [
        ("importe español", "Hotel,F1,2026-08-29,2026-09-28,\"3.630,00\",empresa,0\n", "línea 2: el importe"),
        ("importe a cero", "Hotel,F1,2026-08-29,2026-09-28,0.00,empresa,0\n", "mayor que cero"),
        ("fecha española", "Hotel,F1,29/08/2026,2026-09-28,100.00,empresa,0\n", "`fecha_factura` es '29/08/2026'"),
        ("vence antes de emitirse", "Hotel,F1,2026-08-29,2026-08-01,100.00,empresa,0\n", "es anterior a la factura"),
        ("tipo de cliente inventado", "Hotel,F1,2026-08-29,2026-09-28,100.00,pyme,0\n", "`tipo_cliente` es 'pyme'"),
        ("avisos con texto", "Hotel,F1,2026-08-29,2026-09-28,100.00,empresa,dos\n", "`avisos` es 'dos'"),
        ("sin cliente", ",F1,2026-08-29,2026-09-28,100.00,empresa,0\n", "falta el cliente"),
    ]

    def test_cada_problema_sale_con_su_linea_y_no_se_calcula_nada(self):
        for nombre, linea, mensaje in self.CASOS:
            with self.subTest(nombre):
                codigo, salida, errores = ejecutar(CABECERA + linea, *HOY)
                self.assertEqual(codigo, 1, errores)
                self.assertEqual(salida, "")
                self.assertIn(mensaje, errores)

    def test_la_misma_factura_dos_veces_no_se_cobra_dos_veces(self):
        codigo, _, errores = ejecutar(CABECERA + "Hotel Mirador,F1,2026-08-29,2026-09-28,100.00,empresa,0\nHOTEL MIRADOR ,f1,2026-08-29,2026-09-28,100.00,empresa,0\n", *HOY)
        self.assertEqual(codigo, 1)
        self.assertIn("ya está en la línea 2", errores)

    def test_cabecera_equivocada(self):
        codigo, _, errores = ejecutar("cliente;numero;importe\nHotel;F1;100\n", *HOY)
        self.assertEqual(codigo, 1)
        self.assertIn("La primera línea tiene que llevar estas columnas", errores)

    def test_la_base_no_puede_ser_mayor_que_el_importe_con_iva(self):
        codigo, _, errores = ejecutar(CABECERA.strip() + ",base\nHotel,F1,2026-08-29,2026-09-28,3000.00,empresa,0,3630.00\n", *HOY)
        self.assertEqual(codigo, 1)
        self.assertIn("¿Están cambiadas?", errores)


class Ejemplo(unittest.TestCase):
    """The worked example the skill shows: the figures in its documents are these."""

    ARGUMENTOS = ("--hasta", "2026-10-13", "--plazo-dias", "30")
    CSV = (SKILL / "ejemplos" / "facturas.csv").read_text(encoding="utf-8")

    def test_el_ejemplo_da_lo_que_dice_el_resumen_guardado(self):
        codigo, salida, errores = ejecutar(self.CSV, *self.ARGUMENTOS)
        self.assertEqual(codigo, 0, errores)
        self.assertEqual(json.loads(salida), json.loads((SKILL / "ejemplos" / "resumen.json").read_text(encoding="utf-8")))

    def test_las_cifras_del_ejemplo_hechas_a_mano(self):
        datos = resumen(self.CSV, *self.ARGUMENTOS)
        self.assertEqual(datos["pendiente"], str(D("7260") + D("3630") + D("1815") + D("2900")) + ".00")
        self.assertEqual(datos["intereses_de_demora"], str(D("259.60") + D("15.51") + D("22.75")))
        self.assertEqual(datos["compensacion_por_costes_de_cobro"], "120.00")
        self.assertEqual(datos["total_a_reclamar"], "16022.86")
        self.assertEqual(datos["lo_que_la_ley_anade_al_importe"], "417.86")
        self.assertEqual(datos["facturas_no_vencidas"], 1)

    def test_el_libro_cuadra_con_el_resumen(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl no está instalado (en CI sí lo está).")
        libro = pathlib.Path(tempfile.mkdtemp()) / "reclamacion.xlsx"
        codigo, salida, errores = ejecutar(self.CSV, *self.ARGUMENTOS, "--salida", str(libro))
        self.assertEqual(codigo, 0, errores)
        datos = json.loads(salida)
        hojas = openpyxl.load_workbook(libro)
        self.assertEqual(hojas.sheetnames, ["Reclamaciones", "Por cliente", "Cálculo", "Base legal"])
        filas = list(hojas["Reclamaciones"].iter_rows(min_row=2, values_only=True))
        total = filas[-1]
        self.assertEqual(total[0], "Total")
        self.assertAlmostEqual(total[7], float(datos["total_a_reclamar"]), places=2)
        self.assertAlmostEqual(sum(fila[7] for fila in filas[:-1]), float(datos["total_a_reclamar"]), places=2)
        # The step-by-step sheet adds up to the interest claimed, to the cent per invoice.
        tramos = list(hojas["Cálculo"].iter_rows(min_row=2, values_only=True))
        self.assertAlmostEqual(sum(fila[7] for fila in tramos), float(datos["intereses_de_demora"]), delta=0.02)
        self.assertTrue(all(str(fila[8]).startswith("BOE-A-") for fila in tramos))


if __name__ == "__main__":
    unittest.main()
