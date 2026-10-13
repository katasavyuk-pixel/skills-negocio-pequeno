"""The cash forecast: every figure checked against a sum done by hand."""
import contextlib
import datetime
import io
import json
import pathlib
import sys
import tempfile
import unittest
from decimal import Decimal

RAIZ = pathlib.Path(__file__).resolve().parent.parent
SKILL = RAIZ / "skills" / "prever-caja"
sys.path.insert(0, str(SKILL / "scripts"))

import prever  # noqa: E402

CABECERA = "fecha,concepto,importe,tercero,repite\n"
D = Decimal


def ejecutar(csv, *opciones):
    carpeta = pathlib.Path(tempfile.mkdtemp())
    ruta = carpeta / "movimientos.csv"
    ruta.write_text(csv, encoding="utf-8")
    salida, errores = io.StringIO(), io.StringIO()
    with contextlib.redirect_stdout(salida), contextlib.redirect_stderr(errores):
        codigo = prever.main([str(ruta), *opciones])
    return codigo, salida.getvalue(), errores.getvalue(), carpeta


def resumen(csv, *opciones):
    codigo, salida, errores, _ = ejecutar(csv, *opciones)
    assert codigo == 0, errores
    return json.loads(salida)


def escenario(datos, nombre):
    return next(e for e in datos["escenarios"] if e["escenario"] == nombre)


BASE = ("--saldo", "1000.00", "--fecha", "2026-10-13")


class Semanas(unittest.TestCase):
    def test_la_primera_va_del_dia_del_saldo_al_domingo(self):
        # 13-Oct-2026 is a Tuesday.
        tramos = prever.semanas(datetime.date(2026, 10, 13), 3)
        self.assertEqual(tramos[0], (datetime.date(2026, 10, 13), datetime.date(2026, 10, 18)))
        self.assertEqual(tramos[1], (datetime.date(2026, 10, 19), datetime.date(2026, 10, 25)))
        self.assertEqual(tramos[2][1], datetime.date(2026, 11, 1))

    def test_trece_semanas_por_defecto(self):
        datos = resumen(CABECERA + "2026-10-20,Cobro,100.00,Cliente,no\n2026-10-21,Pago,-50.00,,no\n", *BASE)
        self.assertEqual(datos["semanas"], 13)
        self.assertEqual(datos["hasta"], "2027-01-10")


class Cuentas(unittest.TestCase):
    def test_el_saldo_es_el_de_partida_mas_cobros_menos_pagos(self):
        datos = resumen(CABECERA + "2026-10-20,Cobro,400.00,Cliente,no\n2026-10-27,Pago,-650.50,Proveedor,no\n", *BASE)
        self.assertEqual(D(escenario(datos, "previsto")["saldo_final"]), D("749.50"))

    def test_un_gasto_mensual_se_repite_cada_mes_dentro_del_periodo(self):
        # 31-Oct, 30-Nov and 31-Dec fall inside the 13 weeks; 31-Jan does not.
        datos = resumen(CABECERA + "2026-10-31,Nóminas,-100.00,,mensual\n2026-10-20,Cobro,1.00,Cliente,no\n", *BASE)
        self.assertEqual(D(escenario(datos, "previsto")["saldo_final"]), D("1000.00") + D("1.00") - 3 * D("100.00"))

    def test_el_dia_31_se_ajusta_al_ultimo_dia_de_los_meses_cortos(self):
        self.assertEqual(prever.sumar_meses(datetime.date(2026, 10, 31), 1), datetime.date(2026, 11, 30))
        self.assertEqual(prever.sumar_meses(datetime.date(2026, 10, 31), 4), datetime.date(2027, 2, 28))
        self.assertEqual(prever.sumar_meses(datetime.date(2026, 10, 31), 2), datetime.date(2026, 12, 31))

    def test_un_gasto_mensual_que_empezo_antes_no_cuenta_lo_ya_pagado(self):
        # Rent on the 1st since August: within the period only 1-Nov, 1-Dec and 1-Jan are still to come.
        datos = resumen(CABECERA + "2026-08-01,Alquiler,-100.00,,mensual\n2026-10-20,Cobro,1.00,Cliente,no\n", *BASE)
        self.assertEqual(D(escenario(datos, "previsto")["saldo_final"]), D("1001.00") - 3 * D("100.00"))
        # The August and September rents are in the balance already: they are not unpaid bills.
        self.assertFalse(any("fecha pasada" in aviso for aviso in datos["avisos"]))

    def test_trimestral_son_tres_meses(self):
        datos = resumen(CABECERA + "2026-10-20,IVA modelo 303,-300.00,AEAT,trimestral\n2026-10-21,Cobro,1.00,Cliente,no\n", *BASE, "--semanas", "20")
        # 20-Oct and 20-Jan inside twenty weeks (until 28-Feb-2027).
        self.assertEqual(D(escenario(datos, "previsto")["saldo_final"]), D("1001.00") - 2 * D("300.00"))

    def test_los_decimales_no_se_pierden(self):
        lineas = "".join("2026-10-%02d,Cobro %d,0.10,Cliente %d,no\n" % (14 + i, i, i) for i in range(10))
        datos = resumen(CABECERA + lineas + "2026-10-30,Pago,-0.30,,no\n", "--saldo", "0.20", "--fecha", "2026-10-13")
        self.assertEqual(escenario(datos, "previsto")["saldo_final"], "0.90")


class Escenarios(unittest.TestCase):
    CSV = CABECERA + (
        "2026-10-20,Factura 1,500.00,Grande,no\n"
        "2026-11-10,Factura 2,700.00,Grande,no\n"
        "2026-10-22,Factura 3,300.00,Pequeño,no\n"
        "2026-10-23,Proveedor,-1600.00,Proveedor,no\n"
    )

    def test_con_retraso_solo_se_retrasan_los_cobros(self):
        datos = resumen(self.CSV, *BASE, "--minimo", "500")
        previsto, retraso = escenario(datos, "previsto"), escenario(datos, "con_retraso")
        # Week of 19-Oct as planned: 1000 + 500 + 300 - 1600 = 200. With 14 days of delay nothing has come in yet: -600.
        self.assertEqual(D(previsto["saldo_mas_bajo"]), D("200.00"))
        self.assertEqual(D(retraso["saldo_mas_bajo"]), D("-600.00"))
        self.assertEqual(retraso["primera_semana_en_negativo"], "2026-10-19")
        self.assertIsNone(previsto["primera_semana_en_negativo"])
        self.assertEqual(previsto["primera_semana_bajo_el_minimo"], "2026-10-19")
        # Both end at the same place: the delayed money arrives inside the period.
        self.assertEqual(previsto["saldo_final"], retraso["saldo_final"])

    def test_sin_el_cliente_mayor_desaparecen_todas_sus_facturas(self):
        datos = resumen(self.CSV, *BASE)
        self.assertEqual(datos["cliente_mayor"], {"nombre": "Grande", "importe": "1200.00", "cobros": 2, "parte_de_los_cobros": "80 %"})
        self.assertEqual(D(escenario(datos, "sin_cliente_mayor")["saldo_final"]), D("1000") + D("300") - D("1600"))
        self.assertTrue(any("Grande es el 80 %" in aviso for aviso in datos["avisos"]))

    def test_un_cobro_que_se_retrasa_fuera_del_periodo_no_cuenta(self):
        datos = resumen(CABECERA + "2026-10-20,Factura,500.00,Cliente,no\n2026-10-21,Pago,-100.00,,no\n", *BASE, "--semanas", "2", "--retraso-dias", "30")
        self.assertEqual(D(escenario(datos, "previsto")["saldo_final"]), D("1400.00"))
        self.assertEqual(D(escenario(datos, "con_retraso")["saldo_final"]), D("900.00"))


class Avisos(unittest.TestCase):
    def test_un_cobro_vencido_se_supone_esta_semana_y_se_avisa(self):
        datos = resumen(CABECERA + "2026-09-28,Factura vieja,300.00,Moroso,no\n2026-10-21,Pago,-100.00,,no\n", *BASE)
        self.assertEqual(datos["cobros_vencidos"], "300.00")
        self.assertEqual(D(escenario(datos, "previsto")["saldo_final"]), D("1200.00"))
        self.assertTrue(datos["avisos"][0].startswith("Hay 1 cobro ya vencido por 300,00 €"))

    def test_un_vencimiento_trimestral_sin_pago_de_impuestos_se_pregunta(self):
        sin = resumen(CABECERA + "2026-10-21,Cobro,100.00,Cliente,no\n2026-10-22,Pago,-10.00,,no\n", *BASE)
        self.assertTrue(any("20/10/2026" in aviso and "declaración trimestral" in aviso for aviso in sin["avisos"]))
        con = resumen(CABECERA + "2026-10-21,Cobro,100.00,Cliente,no\n2026-10-20,IVA 3T modelo 303,-10.00,AEAT,no\n", *BASE)
        self.assertFalse(any("declaración trimestral" in aviso for aviso in con["avisos"]))

    def test_sin_cobros_o_sin_pagos_la_prevision_avisa_de_que_esta_coja(self):
        solo_pagos = resumen(CABECERA + "2026-10-20,IVA,-10.00,AEAT,no\n", *BASE)
        self.assertTrue(any("ningún cobro" in aviso for aviso in solo_pagos["avisos"]))
        solo_cobros = resumen(CABECERA + "2026-10-21,Cobro,10.00,Cliente,no\n", *BASE)
        self.assertTrue(any("ningún pago" in aviso for aviso in solo_cobros["avisos"]))


class EntradaMala(unittest.TestCase):
    CASOS = [
        ("importe con coma decimal", "2026-10-20,Cobro,\"1.234,56\",Cliente,no\n", "línea 2: el importe"),
        ("importe con símbolo", "2026-10-20,Cobro,100 €,Cliente,no\n", "línea 2: el importe"),
        ("importe a cero", "2026-10-20,IVA,0.00,AEAT,no\n", "el importe es cero"),
        ("fecha española", "20/10/2026,Cobro,100.00,Cliente,no\n", "línea 2: la fecha"),
        ("repite inventado", "2026-10-20,Cobro,100.00,Cliente,semanal\n", "`repite` es 'semanal'"),
        ("columna de menos", "2026-10-20,Cobro,100.00,no\n", "línea 2: tiene 4 columnas"),
        ("coma sin comillas", "2026-10-20,Cobro, de la obra,100.00,Cliente,no\n", "ponlo entre comillas"),
    ]

    def test_cada_problema_sale_con_su_linea_y_no_se_calcula_nada(self):
        for nombre, linea, mensaje in self.CASOS:
            with self.subTest(nombre):
                codigo, salida, errores, _ = ejecutar(CABECERA + linea, *BASE)
                self.assertEqual(codigo, 1)
                self.assertEqual(salida, "")
                self.assertIn(mensaje, errores)

    def test_se_dicen_todos_los_problemas_de_una_vez(self):
        codigo, _, errores, _ = ejecutar(CABECERA + "20/10/2026,Cobro,100,Cliente,no\n2026-10-21,,x,,no\n", *BASE)
        self.assertEqual(codigo, 1)
        self.assertIn("3 PROBLEMA(S)", errores)

    def test_cabecera_equivocada(self):
        codigo, _, errores, _ = ejecutar("fecha;concepto;importe\n2026-10-20;Cobro;100\n", *BASE)
        self.assertEqual(codigo, 1)
        self.assertIn("La primera línea tiene que ser exactamente", errores)

    def test_fichero_que_no_existe(self):
        salida, errores = io.StringIO(), io.StringIO()
        with contextlib.redirect_stdout(salida), contextlib.redirect_stderr(errores):
            codigo = prever.main(["/no/existe.csv", *BASE])
        self.assertEqual(codigo, 2)
        self.assertIn("No puedo leer", errores.getvalue())


class Ejemplo(unittest.TestCase):
    """The worked example the skill shows: the figures in SKILL.md are these."""

    ARGUMENTOS = ("--saldo", "9420.75", "--fecha", "2026-10-13", "--minimo", "8000")

    def test_el_ejemplo_da_lo_que_dice_el_resumen_guardado(self):
        codigo, salida, errores, _ = ejecutar((SKILL / "ejemplos" / "movimientos.csv").read_text(encoding="utf-8"), *self.ARGUMENTOS)
        self.assertEqual(codigo, 0, errores)
        esperado = json.loads((SKILL / "ejemplos" / "resumen.json").read_text(encoding="utf-8"))
        self.assertEqual(json.loads(salida), esperado)

    def test_la_suma_a_mano_del_ejemplo(self):
        datos = resumen((SKILL / "ejemplos" / "movimientos.csv").read_text(encoding="utf-8"), *self.ARGUMENTOS)
        cobros = D("3630") + D("4840") + D("2900") + D("14520") + D("1815") + D("2400")
        # Three month-ends of payroll and Seguridad Social, three rents, three loan instalments, and the one-offs.
        pagos = 3 * (D("6200") + D("2050")) + 3 * D("1450") + 3 * D("640") + D("3150") + D("3900") + D("1280") + D("6200") + D("980")
        self.assertEqual(D(escenario(datos, "previsto")["saldo_final"]), D("9420.75") + cobros - pagos)

    def test_el_libro_tiene_las_tres_hojas_y_cuadra_con_el_resumen(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl no está instalado (en CI sí lo está).")
        carpeta = pathlib.Path(tempfile.mkdtemp())
        libro = carpeta / "prevision.xlsx"
        codigo, salida, errores, _ = ejecutar((SKILL / "ejemplos" / "movimientos.csv").read_text(encoding="utf-8"), *self.ARGUMENTOS, "--salida", str(libro))
        self.assertEqual(codigo, 0, errores)
        datos = json.loads(salida)
        hojas = openpyxl.load_workbook(libro)
        self.assertEqual(hojas.sheetnames, ["Semanas", "Movimientos", "Avisos"])
        semanas = list(hojas["Semanas"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(semanas), 13)
        self.assertAlmostEqual(semanas[-1][4], float(escenario(datos, "previsto")["saldo_final"]), places=2)
        self.assertAlmostEqual(semanas[-1][6], float(escenario(datos, "sin_cliente_mayor")["saldo_final"]), places=2)
        self.assertIn("Hotel Mirador", hojas["Semanas"]["G1"].value)
        # Every movement of the detail sheet adds up to the change in the balance.
        movimientos = [fila[4] for fila in hojas["Movimientos"].iter_rows(min_row=2, values_only=True)]
        self.assertAlmostEqual(9420.75 + sum(movimientos), semanas[-1][4], places=2)


if __name__ == "__main__":
    unittest.main()
